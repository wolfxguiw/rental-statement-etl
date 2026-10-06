"""Excel report creation without Microsoft Excel automation."""

import os
from pathlib import Path
import tempfile
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .errors import OutputAlreadyExistsError, WorkbookWriteError
from .models import RentalRecord


HEADERS = ("CPF", "NOME", "VALOR DO ALUGUEL")


def write_rental_workbook(
    records: Iterable[RentalRecord], output_path: str | Path, *, overwrite: bool = False
) -> Path:
    """Write the three-column accountant workbook atomically."""

    target = Path(output_path)
    if target.suffix.lower() != ".xlsx":
        raise WorkbookWriteError("o arquivo de saída deve usar a extensão .xlsx")
    if target.exists() and not overwrite:
        raise OutputAlreadyExistsError(f"arquivo de saída já existe: {target}")

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Aluguéis"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for record in records:
        row = sheet.max_row + 1
        cpf_cell = sheet.cell(row=row, column=1, value=record.cpf)
        cpf_cell.number_format = "@"
        cpf_cell.data_type = "s"
        sheet.cell(row=row, column=2, value=record.name).data_type = "s"
        amount_cell = sheet.cell(row=row, column=3, value=float(record.rent))
        amount_cell.number_format = "#,##0.00"

    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 48
    sheet.column_dimensions["C"].width = 22
    sheet.auto_filter.ref = f"A1:C{max(1, sheet.max_row)}"

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{target.stem}.", suffix=".tmp.xlsx", dir=target.parent, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        if target.exists() and not overwrite:
            raise OutputAlreadyExistsError(f"arquivo de saída já existe: {target}")
        os.replace(temporary_path, target)
        return target
    except OutputAlreadyExistsError:
        raise
    except Exception as exc:
        raise WorkbookWriteError(f"falha ao gravar o Excel: {exc}") from exc
    finally:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink(missing_ok=True)
